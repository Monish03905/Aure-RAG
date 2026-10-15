"""Offline document loaders for uploaded knowledge files."""

from __future__ import annotations

import csv
import json
import re
from pathlib import Path
from typing import Dict, List


def clean_text(text: str) -> str:
    text = (text or "").replace("\x00", " ")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def load_txt(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
        try:
            return clean_text(raw.decode(encoding))
        except UnicodeDecodeError:
            continue
    return clean_text(raw.decode("utf-8", errors="ignore"))


def load_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages: List[str] = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    return clean_text("\n\n".join(pages))


def load_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return clean_text("\n".join(parts))


def load_csv(path: Path) -> str:
    rows: List[str] = []
    with path.open("r", encoding="utf-8", errors="ignore", newline="") as f:
        reader = csv.reader(f)
        for i, row in enumerate(reader):
            if i > 2000:
                rows.append("... truncated ...")
                break
            rows.append(" | ".join(cell.strip() for cell in row if cell is not None))
    return clean_text("\n".join(rows))


def load_json_file(path: Path) -> str:
    data = json.loads(path.read_text(encoding="utf-8", errors="ignore"))
    return clean_text(json.dumps(data, ensure_ascii=False, indent=2)[:200000])


LOADERS = {
    ".txt": load_txt,
    ".md": load_txt,
    ".markdown": load_txt,
    ".log": load_txt,
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".csv": load_csv,
    ".json": load_json_file,
}


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    loader = LOADERS.get(suffix)
    if not loader:
        raise ValueError(f"Unsupported file type: {suffix}")
    text = loader(path)
    if not text or len(text) < 20:
        raise ValueError("No usable text extracted from file.")
    return text


def chunk_document(
    text: str,
    *,
    source_name: str,
    doc_id: str,
    chunk_size: int = 700,
    overlap: int = 120,
) -> List[Dict]:
    """Character-window chunks with light sentence awareness."""
    text = clean_text(text)
    if not text:
        return []

    chunks: List[Dict] = []
    start = 0
    index = 0
    length = len(text)

    while start < length:
        end = min(start + chunk_size, length)
        if end < length:
            # Prefer breaking on sentence / paragraph boundaries.
            window = text[start:end]
            break_at = max(
                window.rfind("\n\n"),
                window.rfind(". "),
                window.rfind("? "),
                window.rfind("! "),
            )
            if break_at > chunk_size * 0.45:
                end = start + break_at + 1

        piece = text[start:end].strip()
        if piece:
            chunks.append(
                {
                    "chunk_id": f"{doc_id}_{index}",
                    "text": piece,
                    "answer": piece,
                    "query": source_name,
                    "strategy": "upload_window",
                    "query_id": doc_id,
                    "query_type": "DOCUMENT",
                    "source_lang": "mixed",
                    "target_lang": "mixed",
                    "passage_index": index,
                    "is_selected": True,
                    "chunk_index": index,
                    "source_file": source_name,
                    "doc_id": doc_id,
                    "origin": "upload",
                }
            )
            index += 1

        if end >= length:
            break
        start = max(end - overlap, start + 1)

    return chunks
